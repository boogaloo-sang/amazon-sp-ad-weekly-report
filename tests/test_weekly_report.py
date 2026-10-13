import datetime as dt
import json
from pathlib import Path

from openpyxl import load_workbook

from src.weekly_report import anomalies, build_workbook, reporting_periods


def fixture_payload(config):
    periods = reporting_periods(dt.date(2026, 10, 13))
    period_labels = {key: f"{start}~{end}" for key, (start, end) in periods.items()}
    scopes = ["整体"] + [row["name"] for row in config["owners"]]
    summary = []
    for scope_index, scope in enumerate(scopes):
        for period_index, period in enumerate(("W0", "W1", "W2", "W3", "W4", "Mprev", "Mcur")):
            factor = 1 if scope == "整体" else 0.15 + scope_index * 0.03
            summary.append({
                "scope": scope, "period": period, "period_label": period_labels[period],
                "volume": round((900 + period_index * 20) * factor),
                "net_amount": (70000 + period_index * 1000) * factor,
                "spend": (7000 + period_index * 100) * factor,
                "ad_sales": (42000 + period_index * 800) * factor,
                "ad_orders": round((400 + period_index * 5) * factor),
                "impressions": round((900000 + period_index * 10000) * factor),
                "clicks": round((7000 + period_index * 100) * factor),
            })
    products = []
    for period, factor in (("W3", 1.1), ("W4", 1.0), ("Mprev", 2.8), ("Mcur", 2.6)):
        products.append({
            "country": "US", "sid": "synthetic", "parent_asin": "B0TEST0001", "item_name": "Synthetic product",
            "owners": ["Owner A"], "period": period, "period_label": period_labels[period],
            "volume": round(100 * factor), "net_amount": 5000 * factor, "spend": 250 * factor,
            "ad_sales": 1200 * factor, "ad_orders": round(30 * factor), "impressions": round(10000 * factor), "clicks": round(600 * factor),
        })
    return {"summary": summary, "products": products, "insights": ["Synthetic test only"]}


def test_reporting_periods():
    periods = reporting_periods(dt.date(2026, 10, 13))
    assert periods["W4"] == (dt.date(2026, 10, 5), dt.date(2026, 10, 11))
    assert periods["W1"] == (dt.date(2026, 9, 14), dt.date(2026, 9, 20))
    assert periods["Mcur"] == (dt.date(2026, 10, 1), dt.date(2026, 10, 11))
    assert periods["Mprev"] == (dt.date(2026, 9, 1), dt.date(2026, 9, 11))


def test_anomaly_thresholds():
    rules = {"monthly_previous_sales_min": 10, "weekly_previous_sales_min": 5, "sales_drop_rate": -0.30, "acos_high": 0.35, "acos_increase_pp": 0.10, "tacos_high": 0.20, "tacos_increase_pp": 0.08}
    current = {"volume": 60, "spend": 400, "ad_sales": 800, "net_amount": 1500}
    previous = {"volume": 100, "spend": 100, "ad_sales": 1000, "net_amount": 2000}
    flags = anomalies(current, previous, False, rules)
    assert "周销量异常" in flags and "ACOS异常" in flags and "TACOS异常" in flags


def test_workbook_layout(tmp_path):
    config = json.loads((Path(__file__).parents[1] / "config" / "report.json").read_text(encoding="utf-8"))
    config["owners"] = [
        {"name": "Owner A", "uid": "1001"},
        {"name": "Owner B", "uid": "1002"},
        {"name": "Owner C", "uid": "1003"},
        {"name": "Owner D", "uid": "1004"},
        {"name": "Owner E", "uid": "1005"},
    ]
    periods = reporting_periods(dt.date(2026, 10, 13))
    output = tmp_path / "report.xlsx"
    build_workbook(fixture_payload(config), config, periods, output)
    wb = load_workbook(output, data_only=False)
    assert wb.sheetnames == ["整体概览", "负责人表现", "高花费产品", "口径说明"]
    assert wb["负责人表现"].max_row == 26
    assert wb["高花费产品"]["E7"].value == "A 本周"
    assert wb["高花费产品"]["E10"].value == "D 上月"

